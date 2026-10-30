# -*- coding: utf-8 -*-
"""Gera mapa-conhecimento-mentoria.xlsx a partir de dados_topicos.py.

Layout: tópicos na vertical (uma linha por tópico, à esquerda), mentorados
na horizontal (uma coluna por pessoa, à direita).

Cada mentorado usa sua própria coluna: escreve o nome e o time no topo
(linhas 4-5) e uma nota de 1 a 5 (dropdown) em cada linha de tópico. Isso
permite preenchimento colaborativo simultâneo sem disputa de célula — cada
pessoa mexe só na própria coluna.

As colunas de resumo (D/E/F, à esquerda, ao lado da descrição do tópico)
calculam, por tópico: Média do grupo, % de domínio do grupo e um
sinalizador de gap crítico (média abaixo de 3), agregando a linha inteira
de respostas (todas as colunas de pessoa).

Uso:
    python3 _mentoria/gerar_planilha.py
"""
import os

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter

from dados_topicos import TOPICOS

OUT_PATH = os.path.join(os.path.dirname(__file__), "mapa-conhecimento-mentoria.xlsx")

NUM_COLUNAS_PESSOAS = 40  # colunas em branco pré-formatadas para cadastro (turma inicial: 15)
TITULO = "Mapa de Conhecimento — Mentoria React Native"
LEGENDA_ESCALA = (
    "Escala 1-5: 1 Nunca ouvi falar | 2 Já ouvi falar | "
    "3 Entendo o conceito / uso com apoio | 4 Aplico sozinho | "
    "5 Domino a fundo / referência  —  Cada mentorado usa sua própria coluna: "
    "preencha seu nome e time no topo e uma nota por tópico."
)

COL_TOPICO = 1
COL_DESCRICAO = 2
COL_TRILHAS = 3
COL_MEDIA_GRUPO = 4
COL_PCT_GRUPO = 5
COL_GAP_GRUPO = 6
COL_PRIMEIRA_PESSOA = 7

ROW_TITLE = 1
ROW_LEGEND = 2
ROW_NOME = 4
ROW_TIME = 5
ROW_HEADER = 6
FIRST_TOPIC_ROW = 7

FILL_HEADER = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
FONT_HEADER = Font(color="FFFFFF", bold=True)
FONT_TITLE = Font(size=14, bold=True)
FONT_LEGEND = Font(size=10, italic=True, color="555555")
FONT_LABEL = Font(bold=True, color="1F2937")
FONT_DESC = Font(size=9, color="374151")
FONT_TRILHAS = Font(size=9, italic=True, color="6B7280")
FILL_SUMMARY_LABEL = PatternFill(start_color="E5E7EB", end_color="E5E7EB", fill_type="solid")
FILL_SUMMARY = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
FILL_INPUT = PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid")
BORDER_THIN = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)
FILL_ALT = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")


def build_workbook():
    wb = Workbook()
    ws = wb.active
    ws.title = "Mapa de Conhecimento"

    n_topicos = len(TOPICOS)
    last_topic_row = FIRST_TOPIC_ROW + n_topicos - 1
    row_media_pessoa = last_topic_row + 1

    person_cols = [COL_PRIMEIRA_PESSOA + i for i in range(NUM_COLUNAS_PESSOAS)]
    person_letters = [get_column_letter(c) for c in person_cols]
    last_col = person_cols[-1]
    first_person_letter = person_letters[0]
    last_person_letter = person_letters[-1]

    # Título e legenda
    ws.cell(row=ROW_TITLE, column=1, value=TITULO).font = FONT_TITLE
    ws.cell(row=ROW_LEGEND, column=1, value=LEGENDA_ESCALA).font = FONT_LEGEND
    ws.merge_cells(start_row=ROW_TITLE, start_column=1, end_row=ROW_TITLE, end_column=last_col)
    ws.merge_cells(start_row=ROW_LEGEND, start_column=1, end_row=ROW_LEGEND, end_column=last_col)

    # Linhas Nome / Time no topo de cada coluna de pessoa
    for row, label in ((ROW_NOME, "Nome do mentorado"), (ROW_TIME, "Time")):
        label_cell = ws.cell(row=row, column=1, value=label)
        label_cell.font = FONT_LABEL
        label_cell.fill = FILL_SUMMARY_LABEL
        label_cell.alignment = Alignment(vertical="center")
        label_cell.border = BORDER_THIN
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=COL_PRIMEIRA_PESSOA - 1)
        for col in person_cols:
            cell = ws.cell(row=row, column=col)
            cell.fill = FILL_INPUT
            cell.border = BORDER_THIN
            cell.alignment = Alignment(horizontal="center")

    # Cabeçalho fixo das colunas de descrição/resumo
    header_labels = {
        COL_TOPICO: "Tópico",
        COL_DESCRICAO: "Descrição",
        COL_TRILHAS: "Trilhas",
        COL_MEDIA_GRUPO: "Média do grupo",
        COL_PCT_GRUPO: "% Domínio do grupo",
        COL_GAP_GRUPO: "Gap Crítico do grupo",
    }
    for col, label in header_labels.items():
        c = ws.cell(row=ROW_HEADER, column=col, value=label)
        c.font = FONT_HEADER
        c.fill = FILL_HEADER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER_THIN

    notas_label = ws.cell(row=ROW_HEADER, column=COL_PRIMEIRA_PESSOA, value="Notas dos mentorados (1-5) — uma coluna por pessoa")
    notas_label.font = FONT_HEADER
    notas_label.fill = FILL_HEADER
    notas_label.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.merge_cells(start_row=ROW_HEADER, start_column=COL_PRIMEIRA_PESSOA, end_row=ROW_HEADER, end_column=last_col)
    for col in person_cols:
        ws.cell(row=ROW_HEADER, column=col).border = BORDER_THIN
    ws.row_dimensions[ROW_HEADER].height = 30

    # Uma linha por tópico: descrição visível + fórmulas de resumo do grupo
    for offset, t in enumerate(TOPICOS):
        row = FIRST_TOPIC_ROW + offset
        row_fill = FILL_ALT if offset % 2 == 1 else None

        topico_cell = ws.cell(row=row, column=COL_TOPICO, value=t["topico"])
        topico_cell.font = FONT_LABEL
        topico_cell.alignment = Alignment(vertical="center", wrap_text=True)
        topico_cell.border = BORDER_THIN

        desc_cell = ws.cell(row=row, column=COL_DESCRICAO, value=t["descricao"])
        desc_cell.font = FONT_DESC
        desc_cell.alignment = Alignment(vertical="center", wrap_text=True)
        desc_cell.border = BORDER_THIN

        trilhas_cell = ws.cell(row=row, column=COL_TRILHAS, value=t["trilhas"])
        trilhas_cell.font = FONT_TRILHAS
        trilhas_cell.alignment = Alignment(vertical="center", wrap_text=True)
        trilhas_cell.border = BORDER_THIN

        data_range = f"{first_person_letter}{row}:{last_person_letter}{row}"

        media_cell = ws.cell(row=row, column=COL_MEDIA_GRUPO, value=f'=IFERROR(AVERAGE({data_range}),"")')
        media_cell.number_format = "0.0"
        media_cell.fill = FILL_SUMMARY
        media_cell.border = BORDER_THIN
        media_cell.alignment = Alignment(horizontal="center", vertical="center")

        pct_cell = ws.cell(
            row=row, column=COL_PCT_GRUPO,
            value=f'=IFERROR({get_column_letter(COL_MEDIA_GRUPO)}{row}/5,"")',
        )
        pct_cell.number_format = "0%"
        pct_cell.fill = FILL_SUMMARY
        pct_cell.border = BORDER_THIN
        pct_cell.alignment = Alignment(horizontal="center", vertical="center")

        gap_cell = ws.cell(
            row=row, column=COL_GAP_GRUPO,
            value=(
                f'=IF({get_column_letter(COL_MEDIA_GRUPO)}{row}="","",'
                f'IF({get_column_letter(COL_MEDIA_GRUPO)}{row}<3,"GAP",""))'
            ),
        )
        gap_cell.font = Font(color="B91C1C", bold=True)
        gap_cell.fill = FILL_SUMMARY
        gap_cell.border = BORDER_THIN
        gap_cell.alignment = Alignment(horizontal="center", vertical="center")

        for col in person_cols:
            cell = ws.cell(row=row, column=col)
            cell.border = BORDER_THIN
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if row_fill:
                cell.fill = row_fill

        ws.row_dimensions[row].height = 46

    # Linha de resumo por pessoa (média geral daquele mentorado), logo após os tópicos
    label_cell = ws.cell(row=row_media_pessoa, column=1, value="Média Geral (pessoa)")
    label_cell.font = FONT_LABEL
    label_cell.fill = FILL_SUMMARY_LABEL
    label_cell.alignment = Alignment(vertical="center")
    label_cell.border = BORDER_THIN
    ws.merge_cells(start_row=row_media_pessoa, start_column=1, end_row=row_media_pessoa, end_column=COL_PRIMEIRA_PESSOA - 1)

    for col, letter in zip(person_cols, person_letters):
        cell = ws.cell(
            row=row_media_pessoa, column=col,
            value=f'=IFERROR(AVERAGE({letter}{FIRST_TOPIC_ROW}:{letter}{last_topic_row}),"")',
        )
        cell.number_format = "0.0"
        cell.fill = FILL_SUMMARY
        cell.border = BORDER_THIN
        cell.alignment = Alignment(horizontal="center")

    # Validação de dados: nota 1-5 nas células de resposta
    dv_nota = DataValidation(
        type="list",
        formula1='"1,2,3,4,5"',
        allow_blank=True,
        showDropDown=False,
        showErrorMessage=True,
        errorTitle="Valor inválido",
        error="Escolha um valor inteiro de 1 a 5.",
    )
    ws.add_data_validation(dv_nota)
    dv_nota.add(f"{first_person_letter}{FIRST_TOPIC_ROW}:{last_person_letter}{last_topic_row}")

    # Formatação condicional: escala de cor na coluna "% Domínio do grupo"
    color_scale = ColorScaleRule(
        start_type="min", start_color="F8696B",
        mid_type="percentile", mid_value=50, mid_color="FFEB84",
        end_type="max", end_color="63BE7B",
    )
    ws.conditional_formatting.add(
        f"{get_column_letter(COL_PCT_GRUPO)}{FIRST_TOPIC_ROW}:{get_column_letter(COL_PCT_GRUPO)}{last_topic_row}",
        color_scale,
    )

    # Larguras de coluna
    ws.column_dimensions[get_column_letter(COL_TOPICO)].width = 30
    ws.column_dimensions[get_column_letter(COL_DESCRICAO)].width = 60
    ws.column_dimensions[get_column_letter(COL_TRILHAS)].width = 20
    ws.column_dimensions[get_column_letter(COL_MEDIA_GRUPO)].width = 13
    ws.column_dimensions[get_column_letter(COL_PCT_GRUPO)].width = 15
    ws.column_dimensions[get_column_letter(COL_GAP_GRUPO)].width = 14
    for letter in person_letters:
        ws.column_dimensions[letter].width = 12

    # Congela colunas Tópico/Descrição/Trilhas/resumo (esquerda) e linhas
    # Nome/Time/cabeçalho (topo), para rolar tópicos na vertical e pessoas
    # na horizontal sem perder contexto.
    ws.freeze_panes = ws.cell(row=FIRST_TOPIC_ROW, column=COL_PRIMEIRA_PESSOA)

    return wb


def main():
    wb = build_workbook()
    wb.save(OUT_PATH)
    print(
        f"Escrito: {OUT_PATH} ({len(TOPICOS)} tópicos como linhas, "
        f"{NUM_COLUNAS_PESSOAS} colunas de cadastro pré-formatadas)"
    )


if __name__ == "__main__":
    main()
